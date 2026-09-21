from __future__ import annotations

import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

# La liste de coupe générée par Cutrite (Homag) n'a pas d'en-têtes de colonnes
# exploitables : sa mise en page peut légèrement varier d'une semaine à
# l'autre selon les options d'export. On détecte donc les colonnes utiles par
# leur contenu plutôt que par leur position :
#   - la colonne "clé" (num commande+ligne) contient les mêmes clés que le
#     fichier de lancement (ex. "CVC26070953004") ;
#   - la colonne "matériau" (codes matériaux) contient le code du support
#     (ex. "SAG10HYBR"), ou "PROTECTION" pour les panneaux de protection, qui
#     ne sont pas des pièces client et doivent être exclus, comme dans le
#     fichier de lancement ;
#   - la colonne "grain matching" contient un code de type "T33:2:1" ;
#   - la colonne "type de trace" contient un code de type "P1:TC:..." ou
#     "L0392:TF:...", qui distingue la pièce réellement débitée ("TC") des
#     lignes techniques que Cutrite ajoute pour la même clé commande+ligne
#     (talon/chute "TF", bande de chant "TB", etc.) : ces dernières ne sont
#     pas des pièces et doublonnent le compte si on ne les exclut pas ;
#   - la colonne "quantité" indique combien de pièces identiques une ligne
#     "TC" représente (une même clé peut être débitée en plusieurs fois,
#     avec un nombre de pièces différent à chaque fois selon le nesting) :
#     compter "1 par ligne" est donc faux dès qu'une pièce est nestée en
#     plusieurs lots ; il faut sommer cette colonne, pas les lignes ;
#   - certains exports marquent en plus une ligne technique supplémentaire
#     (réserve/chute) avec le même type de trace "TC" que la vraie pièce ;
#     une autre colonne (souvent un indicateur oui/non) permet de la
#     repérer quand même : on cherche la colonne + valeur dont l'exclusion
#     fait le mieux correspondre les quantités retrouvées aux quantités
#     attendues, en plus du filtrage par type de trace ;
#   - l'identifiant de trace (ex. "P1", "C389", avant ":TC:"/":TF:") est
#     partagé par une pièce (ou sa ligne technique réserve/chute) et sa
#     ligne technique associée : elles devraient avoir la même quantité,
#     ce qui permet un contrôle de cohérence interne de la liste de coupe,
#     indépendant de la comparaison avec le STRAT (voir
#     detecter_incoherences_reserve).
# Une pièce est comptée dès que sa clé commande+ligne est renseignée, que
# son matériau n'est pas "PROTECTION", et que sa ligne est bien une trace
# de type "TC" (quand cette information est disponible dans le fichier),
# sans être par ailleurs marquée par le filtre supplémentaire ci-dessus ;
# sa quantité est celle de la colonne "quantité" si elle a pu être détectée
# avec confiance, sinon 1 par ligne (comportement précédent).
COLONNE_CLE_PAR_DEFAUT = 15  # colonne P
COLONNE_MATERIAU_PAR_DEFAUT = 1  # colonne B
COLONNE_GRAIN_MATCHING_PAR_DEFAUT = 27  # colonne AB

MATERIAUX_EXCLUS = {"PROTECTION"}

MOTIF_GRAIN_MATCHING = re.compile(r"^[A-Za-z]*\d+:\d+:\d+$")
# Groupe 1 : identifiant de trace (ex. "P1", "C389"). Groupe 2 : lettre de
# type ("C" = pièce, "F"/"B" = ligne technique).
MOTIF_TYPE_TRACE = re.compile(r"^([A-Za-z0-9]+):T([A-Z]):")
TYPE_TRACE_PIECE = "C"

# Un modulo ou une bande (ex. "T9/1", "T58/1") n'est jamais une pièce : ce
# sont les panneaux (agglo/strat/sous-face) qui composent ce groupe de
# pièces, qui elles se trouvent ailleurs dans le fichier et se rattachent au
# modulo via leur grain matching ("T9:1 2:1" = fait partie de la bande T9).
# Repérable par sa description (colonne "Noms", en position 0) qui porte
# alors le nom du modulo lui-même — vrai même quand la ligne est par
# ailleurs taguée "TC" comme une vraie pièce.
MOTIF_DESCRIPTION_MODULE = re.compile(r"^T\d+/\d+$")

# La colonne quantité n'est adoptée que si elle explique au moins cette
# proportion des clés du fichier de lancement (sinon, un faux positif sur
# un fichier sans cette colonne compterait n'importe quoi) : mieux vaut
# retomber sur "1 par ligne" que de faire confiance à une colonne qui ne
# correspond pas vraiment aux quantités.
SEUIL_CONFIANCE_COLONNE_QUANTITE = 0.5


class CoupeFormatError(ValueError):
    pass


def _feuille_principale(classeur):
    return max(classeur.worksheets, key=lambda f: f.max_row or 0)


def _charger_lignes(chemin: str | Path) -> list[tuple]:
    chemin = Path(chemin)
    classeur = openpyxl.load_workbook(chemin, data_only=True, read_only=True)
    try:
        feuille = _feuille_principale(classeur)
        lignes = list(feuille.iter_rows(min_row=1, values_only=True))
    finally:
        classeur.close()
    if not lignes:
        raise CoupeFormatError("Le fichier de la liste de coupe est vide.")
    return lignes[1:]


def _detecter_colonne_cle(lignes: list[tuple], cles_strat: set[str]) -> int | None:
    if not cles_strat:
        return None
    nb_colonnes = max((len(ligne) for ligne in lignes), default=0)
    meilleur_index, meilleur_score = None, 0
    for col in range(nb_colonnes):
        score = 0
        for ligne in lignes:
            if col < len(ligne) and ligne[col] is not None and str(ligne[col]).strip() in cles_strat:
                score += 1
        if score > meilleur_score:
            meilleur_index, meilleur_score = col, score
    return meilleur_index


def _detecter_colonne_materiau(lignes: list[tuple]) -> int | None:
    """La colonne matériau est celle qui contient le plus souvent la valeur
    "PROTECTION" (marqueur fiable, présent dans le fichier de lancement comme
    dans la liste de coupe)."""
    nb_colonnes = max((len(ligne) for ligne in lignes), default=0)
    meilleur_index, meilleur_score = None, 0
    for col in range(nb_colonnes):
        score = sum(
            1
            for ligne in lignes
            if col < len(ligne) and ligne[col] is not None and str(ligne[col]).strip().upper() == "PROTECTION"
        )
        if score > meilleur_score:
            meilleur_index, meilleur_score = col, score
    return meilleur_index


def _detecter_colonne_grain_matching(lignes: list[tuple]) -> int | None:
    """La colonne grain matching contient un code du type "T33:2:1"."""
    nb_colonnes = max((len(ligne) for ligne in lignes), default=0)
    meilleur_index, meilleur_score = None, 0
    for col in range(nb_colonnes):
        score = sum(
            1
            for ligne in lignes
            if col < len(ligne) and ligne[col] is not None and MOTIF_GRAIN_MATCHING.match(str(ligne[col]).strip())
        )
        if score > meilleur_score:
            meilleur_index, meilleur_score = col, score
    return meilleur_index


def _detecter_colonne_type_trace(lignes: list[tuple]) -> int | None:
    """La colonne type de trace contient un code du type "P1:TC:..." : la
    lettre après "T" distingue la pièce ("C") des lignes techniques liées
    (talon/chute "F", bande de chant "B"...)."""
    nb_colonnes = max((len(ligne) for ligne in lignes), default=0)
    meilleur_index, meilleur_score = None, 0
    for col in range(nb_colonnes):
        score = sum(
            1
            for ligne in lignes
            if col < len(ligne) and ligne[col] is not None and MOTIF_TYPE_TRACE.match(str(ligne[col]).strip())
        )
        if score > meilleur_score:
            meilleur_index, meilleur_score = col, score
    return meilleur_index


def _est_ligne_piece(ligne: tuple, col_type_trace: int | None) -> bool:
    """Renvoie False pour une ligne technique (talon/chute, bande de chant,
    modulo/bande qui s'auto-référence...) et True pour une vraie pièce.

    Un modulo/bande n'est jamais une pièce, même s'il est par ailleurs tagué
    "TC" comme une vraie pièce (repéré par son nom en colonne 0, voir
    MOTIF_DESCRIPTION_MODULE) : ce contrôle passe avant celui du type de
    trace, qui identifie le reste des lignes techniques (talon/chute, bande
    de chant) quand cette colonne est disponible. Si le type de trace n'a
    pas pu être détecté, les autres lignes restent considérées comme des
    pièces (comportement inchangé)."""
    if len(ligne) > 0 and ligne[0] is not None and MOTIF_DESCRIPTION_MODULE.match(str(ligne[0]).strip().upper()):
        return False

    if col_type_trace is None or col_type_trace >= len(ligne):
        return True
    valeur = ligne[col_type_trace]
    if valeur is None:
        return True
    correspondance = MOTIF_TYPE_TRACE.match(str(valeur).strip())
    if not correspondance:
        return True
    return correspondance.group(2) == TYPE_TRACE_PIECE


def _est_ligne_technique(
    ligne: tuple, col_materiau: int, col_type_trace: int | None, exclusion: tuple[int, str] | None
) -> bool:
    """Renvoie True pour une ligne technique (talon/chute, bande de chant,
    panneau de protection...) : l'inverse des critères de _lignes_pieces(),
    mais sans exiger que la clé soit renseignée (utile pour le contrôle de
    cohérence, qui regroupe justement ces lignes techniques)."""
    materiau = ligne[col_materiau] if col_materiau < len(ligne) else None
    if materiau is not None and str(materiau).strip().upper() in MATERIAUX_EXCLUS:
        return True
    if not _est_ligne_piece(ligne, col_type_trace):
        return True
    if exclusion is not None:
        col_exclu, valeur_exclue = exclusion
        if col_exclu < len(ligne) and ligne[col_exclu] is not None:
            if str(ligne[col_exclu]).strip().upper() == valeur_exclue:
                return True
    return False


def _lignes_pieces(lignes, col_cle, col_materiau, col_type_trace, exclusion=None):
    """Générateur des (clé, ligne) pour les lignes qui représentent une
    pièce réelle : clé renseignée, matériau pas "PROTECTION", ligne de
    type "TC" (quand détectée), et ne correspondant pas au filtre
    supplémentaire `exclusion` = (colonne, valeur), quand détecté."""
    for ligne in lignes:
        if col_cle >= len(ligne) or ligne[col_cle] is None:
            continue
        cle = str(ligne[col_cle]).strip()

        materiau = ligne[col_materiau] if col_materiau < len(ligne) else None
        if materiau is not None and str(materiau).strip().upper() in MATERIAUX_EXCLUS:
            continue

        if not _est_ligne_piece(ligne, col_type_trace):
            continue

        if exclusion is not None:
            col_exclu, valeur_exclue = exclusion
            if col_exclu < len(ligne) and ligne[col_exclu] is not None:
                if str(ligne[col_exclu]).strip().upper() == valeur_exclue:
                    continue

        yield cle, ligne


def _detecter_colonne_quantite(
    lignes: list[tuple],
    col_cle: int,
    col_materiau: int,
    col_type_trace: int | None,
    qte_attendue_par_cle: dict[str, float],
) -> int | None:
    """La colonne quantité est celle dont la somme par clé (sur les lignes
    pièce) reproduit le plus souvent, exactement, la quantité attendue du
    fichier de lancement."""
    if not qte_attendue_par_cle:
        return None
    nb_colonnes = max((len(ligne) for ligne in lignes), default=0)
    meilleur_index, meilleur_score = None, 0
    for col in range(nb_colonnes):
        if col in (col_cle, col_materiau, col_type_trace):
            continue
        sommes: dict[str, float] = defaultdict(float)
        for cle, ligne in _lignes_pieces(lignes, col_cle, col_materiau, col_type_trace):
            valeur = ligne[col] if col < len(ligne) else None
            if valeur is None:
                continue
            try:
                valeur = float(valeur)
            except (TypeError, ValueError):
                continue
            sommes[cle] += valeur
        score = sum(
            1
            for cle, total in sommes.items()
            if cle in qte_attendue_par_cle and abs(total - qte_attendue_par_cle[cle]) < 1e-6
        )
        if score > meilleur_score:
            meilleur_index, meilleur_score = col, score

    seuil = SEUIL_CONFIANCE_COLONNE_QUANTITE * len(qte_attendue_par_cle)
    return meilleur_index if meilleur_score >= seuil else None


def _somme_par_cle(pieces, col_quantite):
    sommes: dict[str, float] = defaultdict(float)
    for cle, ligne in pieces:
        valeur = ligne[col_quantite] if col_quantite is not None and col_quantite < len(ligne) else None
        try:
            quantite = float(valeur) if valeur is not None else 1.0
        except (TypeError, ValueError):
            quantite = 1.0
        sommes[cle] += quantite
    return sommes


def _detecter_exclusion_supplementaire(
    lignes: list[tuple],
    col_cle: int,
    col_materiau: int,
    col_type_trace: int | None,
    col_quantite: int | None,
    qte_attendue_par_cle: dict[str, float],
) -> tuple[int, str] | None:
    """Cherche une colonne + valeur dont l'exclusion, en plus du filtrage
    par type de trace, améliore encore la correspondance entre quantités
    retrouvées et quantités attendues (cas d'une ligne technique marquée
    par erreur du même type "TC" que la vraie pièce, mais repérable par un
    indicateur dans une autre colonne, ex. oui/non)."""
    if not qte_attendue_par_cle:
        return None

    pieces = list(_lignes_pieces(lignes, col_cle, col_materiau, col_type_trace))

    def score(sommes: dict[str, float]) -> int:
        return sum(
            1
            for cle, total in sommes.items()
            if cle in qte_attendue_par_cle and abs(total - qte_attendue_par_cle[cle]) < 1e-6
        )

    score_base = score(_somme_par_cle(pieces, col_quantite))

    nb_colonnes = max((len(ligne) for _, ligne in pieces), default=0)
    meilleur, meilleur_score = None, score_base
    for col in range(nb_colonnes):
        if col in (col_cle, col_materiau, col_type_trace, col_quantite):
            continue
        valeurs: set[str] = set()
        for _cle, ligne in pieces:
            if col < len(ligne) and ligne[col] is not None:
                valeurs.add(str(ligne[col]).strip().upper())
        valeurs.discard("")
        # On se limite à des colonnes catégorielles (peu de valeurs
        # distinctes) : c'est le profil attendu d'un indicateur, et ça
        # évite de tester des colonnes de texte libre sans rapport.
        if not (2 <= len(valeurs) <= 6):
            continue
        for valeur in valeurs:
            pieces_filtrees = [
                (cle, ligne)
                for cle, ligne in pieces
                if not (col < len(ligne) and ligne[col] is not None and str(ligne[col]).strip().upper() == valeur)
            ]
            s = score(_somme_par_cle(pieces_filtrees, col_quantite))
            if s > meilleur_score:
                meilleur, meilleur_score = (col, valeur), s

    seuil = SEUIL_CONFIANCE_COLONNE_QUANTITE * len(qte_attendue_par_cle)
    return meilleur if meilleur is not None and meilleur_score >= seuil else None


def detecter_incoherences_reserve(
    donnees: list[tuple],
    col_cle: int,
    col_materiau: int,
    col_type_trace: int | None,
    exclusion: tuple[int, str] | None,
    col_quantite: int | None,
    col_grain: int | None,
) -> list[dict]:
    """Contrôle de cohérence interne de la liste de coupe, indépendant de
    la comparaison avec le STRAT : les lignes techniques (réserve/chute...)
    partageant le même identifiant de trace (ex. "C389") devraient avoir
    la même quantité. Si ce n'est pas le cas, c'est le signe d'une
    modification ou d'une incohérence dans le fichier de coupe lui-même —
    même si ces lignes n'entrent pas dans le compte des pièces réelles.
    Chaque incohérence rapporte aussi la description (colonne "Noms"/
    "Description") et le grain matching des lignes concernées, pour les
    identifier facilement dans le fichier."""
    if col_type_trace is None:
        return []

    groupes: dict[tuple[str, str], dict] = {}
    for ligne in donnees:
        if col_cle >= len(ligne) or ligne[col_cle] is None:
            continue
        if not _est_ligne_technique(ligne, col_materiau, col_type_trace, exclusion):
            continue
        if col_type_trace >= len(ligne) or ligne[col_type_trace] is None:
            continue

        correspondance = MOTIF_TYPE_TRACE.match(str(ligne[col_type_trace]).strip())
        if not correspondance:
            continue

        cle = str(ligne[col_cle]).strip()
        prefixe = correspondance.group(1)
        valeur = ligne[col_quantite] if col_quantite is not None and col_quantite < len(ligne) else None
        try:
            quantite = float(valeur) if valeur is not None else None
        except (TypeError, ValueError):
            quantite = None

        description = str(ligne[0]).strip() if len(ligne) > 0 and ligne[0] is not None else None
        grain = str(ligne[col_grain]).strip() if col_grain is not None and col_grain < len(ligne) and ligne[col_grain] is not None else None

        clef = (cle, prefixe)
        groupe = groupes.setdefault(
            clef, {"cle": cle, "id_trace": prefixe, "quantites": [], "descriptions": set(), "grains": set()}
        )
        if quantite is not None:
            groupe["quantites"].append(quantite)
        if description:
            groupe["descriptions"].add(description)
        if grain:
            groupe["grains"].add(grain)

    incoherences = []
    for groupe in groupes.values():
        distinctes = sorted(set(groupe["quantites"]))
        if len(distinctes) <= 1:
            continue  # quantité inhabituelle mais cohérente entre les deux lignes de sa paire : pas une anomalie en soi
        incoherences.append(
            {
                "cle": groupe["cle"],
                "id_trace": groupe["id_trace"],
                "description": " / ".join(sorted(groupe["descriptions"])) or None,
                "grain": " / ".join(sorted(groupe["grains"])) or None,
                "type": "ecart",
                "quantites": distinctes,
            }
        )
    return incoherences


# Le grain matching d'une pièce liste les positions qu'elle occupe dans
# son module/bande de calage de fil (ex. "T76:1 2 3:1" = positions
# 1, 2, 3). Le nombre de positions listées doit correspondre à la
# quantité réellement trouvée pour cette pièce : sur un fichier réel
# observé, 321 clés sur 323 concordent exactement — un écart est donc un
# signal fiable, même quand le total global reste correct par ailleurs
# (des pièces peuvent manquer dans un module sans casser le compte
# global, qui raisonne par clé et pas par position).
MOTIF_POSITIONS_GRAIN = re.compile(r"^[A-Za-z]*\d+:([\d ]+):\d+$")


def detecter_ecarts_grain_matching(
    pieces: list[tuple[str, tuple]], col_grain: int | None, compteur: Counter
) -> list[dict]:
    if col_grain is None:
        return []

    par_cle: dict[str, dict] = {}
    for cle, ligne in pieces:
        if col_grain >= len(ligne) or ligne[col_grain] is None:
            continue
        brut = str(ligne[col_grain]).strip()
        correspondance = MOTIF_POSITIONS_GRAIN.match(brut)
        if not correspondance:
            continue
        infos = par_cle.setdefault(cle, {"positions": set(), "bruts": set()})
        infos["positions"].update(correspondance.group(1).split())
        infos["bruts"].add(brut)

    ecarts = []
    for cle, infos in par_cle.items():
        nb_positions = len(infos["positions"])
        quantite_trouvee = compteur.get(cle, 0)
        if abs(nb_positions - quantite_trouvee) > 1e-6:
            ecarts.append(
                {
                    "type": "grain_matching",
                    "cle": cle,
                    "nb_positions": nb_positions,
                    "quantite_trouvee": quantite_trouvee,
                    "grain": " / ".join(sorted(infos["bruts"])),
                }
            )
    return ecarts


def _compter_unitaires(pieces: list[tuple[str, tuple]], col_grain: int | None, compteur: Counter) -> int | None:
    """Compte les clés trouvées qui n'ont aucune valeur de grain matching
    sur aucune de leurs lignes ("unitaires" : pas de calage de fil
    nécessaire). Renvoie None si la colonne grain matching n'a pas pu être
    détectée."""
    if col_grain is None:
        return None
    a_grain: set[str] = set()
    for cle, ligne in pieces:
        if col_grain < len(ligne) and ligne[col_grain] is not None and str(ligne[col_grain]).strip():
            a_grain.add(cle)
    return sum(1 for cle in compteur if cle not in a_grain)


def lire_coupe(
    chemin: str | Path, cles_strat: set[str], qte_attendue_par_cle: dict[str, float] | None = None
) -> tuple[Counter, list[dict], int | None]:
    """Compte, pour chaque clé commande+ligne, le nombre de pièces dans la
    liste de coupe Cutrite (panneaux de protection et lignes techniques
    exclus). Si la colonne "quantité" par ligne peut être détectée avec
    confiance (via `qte_attendue_par_cle`), une même clé nestée en
    plusieurs lots est comptée correctement au lieu d'être sous-évaluée
    (1 par ligne, quel que soit le nombre de pièces qu'elle représente).

    Renvoie aussi les incohérences détectées entre lignes techniques (voir
    detecter_incoherences_reserve) : ces dernières peuvent exister même
    quand le compteur retourné est parfaitement cohérent avec le STRAT."""

    donnees = _charger_lignes(chemin)

    col_cle = _detecter_colonne_cle(donnees, cles_strat)
    if col_cle is None:
        col_cle = COLONNE_CLE_PAR_DEFAUT

    col_materiau = _detecter_colonne_materiau(donnees)
    if col_materiau is None:
        col_materiau = COLONNE_MATERIAU_PAR_DEFAUT

    col_type_trace = _detecter_colonne_type_trace(donnees)
    col_grain = _detecter_colonne_grain_matching(donnees)

    col_quantite = _detecter_colonne_quantite(
        donnees, col_cle, col_materiau, col_type_trace, qte_attendue_par_cle or {}
    )
    exclusion = _detecter_exclusion_supplementaire(
        donnees, col_cle, col_materiau, col_type_trace, col_quantite, qte_attendue_par_cle or {}
    )

    pieces = list(_lignes_pieces(donnees, col_cle, col_materiau, col_type_trace, exclusion))

    compteur: Counter = Counter()
    for cle, ligne in pieces:
        quantite = 1.0
        if col_quantite is not None and col_quantite < len(ligne) and ligne[col_quantite] is not None:
            try:
                quantite = float(ligne[col_quantite])
            except (TypeError, ValueError):
                quantite = 1.0
        compteur[cle] += quantite

    if sum(compteur.values()) == 0:
        raise CoupeFormatError(
            "Aucune pièce n'a été reconnue dans la liste de coupe : le format "
            "du fichier semble différent de celui attendu."
        )

    incoherences_reserve = detecter_incoherences_reserve(
        donnees, col_cle, col_materiau, col_type_trace, exclusion, col_quantite, col_grain
    )
    ecarts_grain_matching = detecter_ecarts_grain_matching(pieces, col_grain, compteur)

    nb_unitaires = _compter_unitaires(pieces, col_grain, compteur)

    return compteur, incoherences_reserve + ecarts_grain_matching, nb_unitaires


def lire_grain_matching(chemin: str | Path, cles_strat: set[str] | None = None) -> dict[str, list[str]]:
    """Renvoie, pour chaque clé commande+ligne trouvée dans la liste de
    coupe, la ou les valeurs de grain matching associées (ex. "T33:2:1")."""

    donnees = _charger_lignes(chemin)

    col_cle = _detecter_colonne_cle(donnees, cles_strat or set())
    if col_cle is None:
        col_cle = COLONNE_CLE_PAR_DEFAUT

    col_grain = _detecter_colonne_grain_matching(donnees)
    if col_grain is None:
        col_grain = COLONNE_GRAIN_MATCHING_PAR_DEFAUT

    col_type_trace = _detecter_colonne_type_trace(donnees)
    col_materiau = _detecter_colonne_materiau(donnees)
    if col_materiau is None:
        col_materiau = COLONNE_MATERIAU_PAR_DEFAUT

    resultat: dict[str, list[str]] = defaultdict(list)
    for cle, ligne in _lignes_pieces(donnees, col_cle, col_materiau, col_type_trace):
        # On enregistre la clé même sans grain matching, pour distinguer une
        # pièce unitaire (clé trouvée, pas de grain matching) d'une clé
        # absente de la liste de coupe.
        resultat.setdefault(cle, [])

        valeur = ligne[col_grain] if col_grain < len(ligne) else None
        if valeur is None:
            continue
        valeur = str(valeur).strip()
        if valeur and valeur not in resultat[cle]:
            resultat[cle].append(valeur)

    return dict(resultat)
