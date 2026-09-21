from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import openpyxl

from .models import LigneStrat

# En-têtes attendus dans le fichier de lancement (export STRAT). Les noms de
# colonnes sont stables d'une semaine à l'autre car ils viennent de l'ERP.
COL_COMMANDE = "conf_commande"
COL_LIGNE = "conf_ligne"
COL_ARTICLE = "conf_article"
COL_QTE = "comdet_qte"
COL_LANCEMENT = "comprod_lancement"
COL_CLE = "commande+ligne"
COL_AGGLO = "agglo"
COL_LIVSEM = "com_livsem"
COL_LIVAN = "com_livan"

REQUIRED_HEADERS = (COL_COMMANDE, COL_LIGNE, COL_ARTICLE, COL_QTE)

# Un panneau de protection (support = "PROTECTION") n'est pas une pièce
# commandée par le client : il ne doit pas être compté dans le contrôle.
MATERIAUX_EXCLUS = {"PROTECTION"}


class StratFormatError(ValueError):
    pass


def _trouver_feuille_et_entetes(chemin: Path):
    classeur = openpyxl.load_workbook(chemin, data_only=True, read_only=True)
    for feuille in classeur.worksheets:
        ligne_entetes = next(feuille.iter_rows(min_row=1, max_row=1), None)
        if ligne_entetes is None:
            continue
        entetes = {
            str(cellule.value).strip(): index
            for index, cellule in enumerate(ligne_entetes, start=1)
            if cellule.value is not None
        }
        if all(entete in entetes for entete in REQUIRED_HEADERS):
            return classeur, feuille, entetes
    classeur.close()
    raise StratFormatError(
        "Aucune feuille du fichier de lancement ne contient les colonnes "
        f"attendues ({', '.join(REQUIRED_HEADERS)})."
    )


def _cle_ligne(commande: str, ligne) -> str:
    try:
        numero = int(float(ligne))
    except (TypeError, ValueError):
        return f"{commande}{ligne}"
    return f"{commande}{numero:03d}"


def _fmt_ligne(ligne) -> str:
    if ligne is None:
        return ""
    try:
        return str(int(float(ligne)))
    except (TypeError, ValueError):
        return str(ligne)


def lire_strat(chemin: str | Path) -> tuple[list[LigneStrat], list[dict], list[dict]]:
    """Lit le fichier de lancement et agrège les lignes par clé
    commande+ligne (les quantités des lignes partageant une même clé sont
    additionnées).

    Renvoie aussi :
    - les doublons détectés : une même clé commande+ligne apparaissant sur
      plusieurs lignes (hors panneaux "PROTECTION", exclus plus haut et qui
      ne sont donc jamais la cause d'un doublon signalé ici) peut indiquer
      une pièce lancée deux fois par erreur dans l'ERP ;
    - les semaines de livraison distinctes trouvées (colonne "com_livsem") :
      un fichier de lancement ne devrait normalement en contenir qu'une
      seule ; plusieurs semaines différentes peuvent signaler des pièces
      mélangées par erreur entre deux lancements."""

    chemin = Path(chemin)
    classeur, feuille, entetes = _trouver_feuille_et_entetes(chemin)
    try:
        idx_commande = entetes[COL_COMMANDE]
        idx_ligne = entetes[COL_LIGNE]
        idx_article = entetes[COL_ARTICLE]
        idx_qte = entetes[COL_QTE]
        idx_lancement = entetes.get(COL_LANCEMENT)
        idx_cle = entetes.get(COL_CLE)
        idx_agglo = entetes.get(COL_AGGLO)
        idx_livsem = entetes.get(COL_LIVSEM)
        idx_livan = entetes.get(COL_LIVAN)

        qte_par_cle: dict[str, float] = defaultdict(float)
        infos_par_cle: dict[str, tuple[str, str, str, str]] = {}
        lignes_brutes_par_cle: dict[str, list[tuple[str, float]]] = defaultdict(list)
        compte_par_semaine: dict[tuple[str, str], int] = defaultdict(int)

        nb_colonnes = feuille.max_column

        def valeur(ligne_donnees, index: int | None):
            if not index or index > len(ligne_donnees):
                return None
            return ligne_donnees[index - 1].value

        for ligne_donnees in feuille.iter_rows(min_row=2, max_col=nb_colonnes):
            commande = valeur(ligne_donnees, idx_commande)
            if commande is None:
                continue
            commande = str(commande).strip()

            agglo = valeur(ligne_donnees, idx_agglo)
            if agglo is not None and str(agglo).strip().upper() in MATERIAUX_EXCLUS:
                continue
            numero_ligne = valeur(ligne_donnees, idx_ligne)
            article = valeur(ligne_donnees, idx_article) or ""
            qte = valeur(ligne_donnees, idx_qte) or 0
            lancement = valeur(ligne_donnees, idx_lancement) or ""

            cle = valeur(ligne_donnees, idx_cle)
            cle = str(cle).strip() if cle else _cle_ligne(commande, numero_ligne)

            qte_flottante = float(qte)
            qte_par_cle[cle] += qte_flottante
            infos_par_cle[cle] = (
                commande,
                _fmt_ligne(numero_ligne),
                str(article),
                str(lancement) if lancement else "",
            )
            lignes_brutes_par_cle[cle].append((str(article), qte_flottante))

            livsem = valeur(ligne_donnees, idx_livsem)
            if livsem is not None and str(livsem).strip() != "":
                livan = valeur(ligne_donnees, idx_livan)
                compte_par_semaine[(str(livan) if livan is not None else "", str(livsem))] += 1

        resultat = []
        for cle, qte in qte_par_cle.items():
            commande, numero_ligne, article, lancement = infos_par_cle[cle]
            resultat.append(
                LigneStrat(
                    cle=cle,
                    commande=commande,
                    ligne=numero_ligne,
                    article=article,
                    lancement=lancement,
                    qte_attendue=qte,
                )
            )

        doublons = []
        for cle, occurrences in lignes_brutes_par_cle.items():
            if len(occurrences) <= 1:
                continue
            commande, numero_ligne, _, _ = infos_par_cle[cle]
            doublons.append(
                {
                    "cle": cle,
                    "commande": commande,
                    "ligne": numero_ligne,
                    "nb_lignes": len(occurrences),
                    "articles": [a for a, _ in occurrences],
                    "quantites": [q for _, q in occurrences],
                }
            )

        semaines = [
            {"annee": annee, "semaine": semaine, "nb_pieces": nb}
            for (annee, semaine), nb in sorted(compte_par_semaine.items())
        ]

        return resultat, doublons, semaines
    finally:
        classeur.close()
