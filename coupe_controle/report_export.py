from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import RapportComparaison, Statut

EN_TETES = [
    "Commande",
    "Ligne",
    "Article",
    "Lancement",
    "Qté attendue (STRAT)",
    "Qté trouvée (coupe)",
    "Écart",
    "Statut",
]

REMPLISSAGE_PAR_STATUT = {
    Statut.OK: PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    Statut.MANQUANT: PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    Statut.EN_TROP: PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    Statut.ECART_QTE: PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
}


def exporter_rapport(rapport: RapportComparaison, chemin: str | Path) -> None:
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Contrôle coupe"

    feuille.append([
        "Total attendu",
        rapport.total_attendu,
        "Total trouvé",
        rapport.total_trouve,
        "Anomalies",
        rapport.nb_anomalies,
    ])
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    feuille.append([])

    ligne_entetes = feuille.max_row + 1
    feuille.append(EN_TETES)
    for cellule in feuille[ligne_entetes]:
        cellule.font = Font(bold=True)
        cellule.alignment = Alignment(horizontal="center")

    for ligne in rapport.lignes:
        feuille.append(
            [
                ligne.commande,
                ligne.ligne,
                ligne.article,
                ligne.lancement,
                ligne.qte_attendue,
                ligne.qte_trouvee,
                ligne.ecart,
                ligne.statut.value,
            ]
        )
        remplissage = REMPLISSAGE_PAR_STATUT.get(ligne.statut)
        if remplissage is not None:
            for cellule in feuille[feuille.max_row]:
                cellule.fill = remplissage

    largeurs = [16, 8, 22, 14, 20, 20, 8, 20]
    for index, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[feuille.cell(row=1, column=index).column_letter].width = largeur

    classeur.save(chemin)
